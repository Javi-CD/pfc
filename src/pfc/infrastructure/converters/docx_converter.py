from pathlib import Path

from docx import Document
from docx2pdf import convert  # type: ignore

from pfc.domain.exceptions import ConversionError
from pfc.domain.models import ConversionResult


class DocxToTxtConverter:
    def supports(self, source_format: str, target_format: str) -> bool:
        return source_format.lower() == "docx" and target_format.lower() == "txt"

    def convert(self, source: Path, destination: Path) -> ConversionResult:
        size_before = source.stat().st_size
        try:
            doc = Document(str(source))
            text_content = [p.text for p in doc.paragraphs]

            with open(destination, "w", encoding="utf-8") as f:
                f.write("\n".join(text_content))
        except Exception as e:
            raise ConversionError(f"Failed to extract text from DOCX: {e}")

        size_after = destination.stat().st_size
        return ConversionResult(
            source=source,
            destination=destination,
            source_format="docx",
            target_format="txt",
            size_before=size_before,
            size_after=size_after,
        )


class TxtToDocxConverter:
    def supports(self, source_format: str, target_format: str) -> bool:
        return source_format.lower() == "txt" and target_format.lower() == "docx"

    def convert(self, source: Path, destination: Path) -> ConversionResult:
        size_before = source.stat().st_size
        try:
            with open(source, "r", encoding="utf-8") as f:
                lines = f.readlines()

            doc = Document()
            for line in lines:
                doc.add_paragraph(line.rstrip("\n"))
            doc.save(str(destination))
        except Exception as e:
            raise ConversionError(f"Failed to create DOCX from TXT: {e}")

        size_after = destination.stat().st_size
        return ConversionResult(
            source=source,
            destination=destination,
            source_format="txt",
            target_format="docx",
            size_before=size_before,
            size_after=size_after,
        )


class DocxToPdfConverter:
    def supports(self, source_format: str, target_format: str) -> bool:
        return source_format.lower() == "docx" and target_format.lower() == "pdf"

    def convert(self, source: Path, destination: Path) -> ConversionResult:
        size_before = source.stat().st_size
        try:
            convert(str(source), str(destination))
        except Exception as e:
            raise ConversionError(f"Failed to convert DOCX to PDF: {e}")

        if not destination.exists():
            raise ConversionError(
                "DOCX to PDF conversion silently failed (Is Microsoft Word installed?)"
            )

        size_after = destination.stat().st_size
        return ConversionResult(
            source=source,
            destination=destination,
            source_format="docx",
            target_format="pdf",
            size_before=size_before,
            size_after=size_after,
        )
