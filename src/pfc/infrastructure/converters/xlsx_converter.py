import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook

from pfc.domain.exceptions import ConversionError
from pfc.domain.models import ConversionResult


class CsvToXlsxConverter:
    def supports(self, source_format: str, target_format: str) -> bool:
        return source_format.lower() == "csv" and target_format.lower() == "xlsx"

    def convert(self, source: Path, destination: Path) -> ConversionResult:
        size_before = source.stat().st_size

        wb = Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet()

        with open(source, "r", encoding="utf-8") as f:
            reader = csv.reader(f)
            for row in reader:
                ws.append(row)

        wb.save(destination)

        size_after = destination.stat().st_size

        return ConversionResult(
            source=source,
            destination=destination,
            source_format="csv",
            target_format="xlsx",
            size_before=size_before,
            size_after=size_after,
        )


class XlsxToCsvConverter:
    def supports(self, source_format: str, target_format: str) -> bool:
        return source_format.lower() == "xlsx" and target_format.lower() == "csv"

    def convert(self, source: Path, destination: Path) -> ConversionResult:
        size_before = source.stat().st_size

        try:
            wb = load_workbook(filename=source, data_only=True)
            ws = wb.active
            if ws is None:
                raise ConversionError("XLSX file has no active sheet")
        except Exception as e:
            raise ConversionError(f"Failed to read XLSX file: {e}")

        with open(destination, "w", encoding="utf-8", newline="") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                # Convert None to empty string
                row_str = ["" if cell is None else str(cell) for cell in row]
                writer.writerow(row_str)

        size_after = destination.stat().st_size

        return ConversionResult(
            source=source,
            destination=destination,
            source_format="xlsx",
            target_format="csv",
            size_before=size_before,
            size_after=size_after,
        )
