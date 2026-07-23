from pathlib import Path

from docx import Document

from pfc.infrastructure.converters.docx_converter import (
    DocxToPdfConverter,
    DocxToTxtConverter,
    TxtToDocxConverter,
)


def test_docx_to_txt(tmp_path: Path):
    docx_file = tmp_path / "test.docx"
    doc = Document()
    doc.add_paragraph("Hello")
    doc.save(str(docx_file))

    txt_file = tmp_path / "test.txt"
    converter = DocxToTxtConverter()
    assert converter.supports("docx", "txt")
    converter.convert(docx_file, txt_file)
    assert txt_file.read_text(encoding="utf-8") == "Hello"

def test_txt_to_docx(tmp_path: Path):
    txt_file = tmp_path / "test.txt"
    txt_file.write_text("World", encoding="utf-8")
    docx_file = tmp_path / "test.docx"

    converter = TxtToDocxConverter()
    assert converter.supports("txt", "docx")
    converter.convert(txt_file, docx_file)

    doc = Document(str(docx_file))
    assert doc.paragraphs[0].text == "World"

def test_docx_to_pdf_supports():
    converter = DocxToPdfConverter()
    assert converter.supports("docx", "pdf")
    assert not converter.supports("pdf", "docx")
