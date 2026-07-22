import csv
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from pfc.domain.exceptions import ConversionError
from pfc.infrastructure.converters.xlsx_converter import CsvToXlsxConverter, XlsxToCsvConverter


def test_csv_to_xlsx_converter(tmp_path: Path) -> None:
    csv_file = tmp_path / "test.csv"
    csv_file.write_text("name,age\nAlice,30\nBob,25", encoding="utf-8")
    xlsx_file = tmp_path / "test.xlsx"

    converter = CsvToXlsxConverter()
    assert converter.supports("csv", "xlsx")
    assert not converter.supports("xlsx", "csv")

    result = converter.convert(csv_file, xlsx_file)

    assert result.source_format == "csv"
    assert result.target_format == "xlsx"

    # Verify the output
    wb = load_workbook(xlsx_file, data_only=True)
    ws = wb.active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3
    assert rows[0] == ("name", "age")
    assert rows[1] == ("Alice", "30")


def test_xlsx_to_csv_converter(tmp_path: Path) -> None:
    xlsx_file = tmp_path / "test.xlsx"
    wb = Workbook()
    ws = wb.active
    if ws is not None:
        ws.append(["name", "age"])
        ws.append(["Alice", 30])
        ws.append(["Bob", 25])
    wb.save(xlsx_file)

    csv_file = tmp_path / "test.csv"

    converter = XlsxToCsvConverter()
    assert converter.supports("xlsx", "csv")
    assert not converter.supports("csv", "xlsx")

    result = converter.convert(xlsx_file, csv_file)

    assert result.source_format == "xlsx"
    assert result.target_format == "csv"

    # Verify the output
    with open(csv_file, "r", encoding="utf-8") as f:
        reader = list(csv.reader(f))
    assert len(reader) == 3
    assert reader[0] == ["name", "age"]
    assert reader[1] == ["Alice", "30"]


def test_xlsx_to_csv_invalid_file(tmp_path: Path) -> None:
    invalid_file = tmp_path / "invalid.xlsx"
    invalid_file.write_text("not an xlsx file", encoding="utf-8")
    csv_file = tmp_path / "test.csv"

    converter = XlsxToCsvConverter()
    with pytest.raises(ConversionError):
        converter.convert(invalid_file, csv_file)
